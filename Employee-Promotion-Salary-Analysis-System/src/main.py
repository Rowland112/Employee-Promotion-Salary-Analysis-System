"""
Employee Promotion & Salary Analysis System

A Python business automation application that:
- Loads employee records from Excel
- Validates employee IDs and email addresses
- Evaluates promotion eligibility
- Performs salary-position analysis using a configurable benchmark
- Generates Excel, TXT, and DOCX reports
"""

from pathlib import Path
from datetime import datetime
import re

import openpyxl
from docx import Document


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
REPORTS_DIR = PROJECT_ROOT / "reports"
INPUT_FILE = DATA_DIR / "EmployeeData.xlsx"

EMPLOYEE_ID_PATTERN = r"^EMP\s\d{3}$"
EMAIL_PATTERN = r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$"

# This is a demo benchmark used for the portfolio project.
# It can be replaced later with an approved external salary data source.
BASE_MARKET_SALARY = 7000
POSITION_BENCHMARK_INCREMENT = 100


def validate_employee_id(employee_id):
    """Return True when the employee ID matches the expected format."""
    return bool(re.match(EMPLOYEE_ID_PATTERN, str(employee_id).strip()))


def validate_email(email):
    """Return True when the email has a basic valid format."""
    return bool(re.match(EMAIL_PATTERN, str(email).strip()))


def extract_id_number(employee_id):
    """Return the numeric portion of an employee ID."""
    return str(employee_id).replace(" ", "")[3:]


def load_employee_sheet(file_path=INPUT_FILE):
    """Load the first worksheet from the employee Excel file."""
    if not file_path.exists():
        raise FileNotFoundError(
            "EmployeeData.xlsx was not found in the data folder."
        )

    workbook = openpyxl.load_workbook(file_path)
    return workbook.active


def check_promotion(years, score, attendance, warnings, training):
    """Apply the project's promotion eligibility rules."""
    if (
        years >= 3
        and score >= 65
        and attendance >= 80
        and warnings <= 3
        and str(training).strip().lower() == "completed"
    ):
        return "Eligible"

    return "Not Eligible"


def get_market_salary(position):
    """
    Return a configurable demonstration salary benchmark.

    The original college project attempted web scraping. For a reliable
    portfolio application, this version avoids pretending that a scraped
    number is an authoritative market salary. The benchmark can later be
    connected to a verified salary-data source.
    """
    position_length = len(str(position).strip())
    return BASE_MARKET_SALARY + (
        position_length * POSITION_BENCHMARK_INCREMENT
    )


def display_employees(sheet):
    """Display a simple employee directory in the terminal."""
    print("\nEMPLOYEE LIST")
    print("=" * 70)

    for row in range(2, sheet.max_row + 1):
        employee_id = sheet.cell(row, 1).value
        first_name = sheet.cell(row, 2).value
        last_name = sheet.cell(row, 3).value
        department = sheet.cell(row, 5).value

        print(
            str(employee_id)
            + " - "
            + str(first_name)
            + " "
            + str(last_name)
            + " - "
            + str(department)
        )


def search_employee(sheet):
    """Search for an employee by ID."""
    search_id = input("Enter Employee ID (example: EMP 101): ").strip()
    found = False

    for row in range(2, sheet.max_row + 1):
        employee_id = str(sheet.cell(row, 1).value).strip()

        if employee_id.upper() == search_id.upper():
            print("\nEMPLOYEE FOUND")
            print("=" * 40)

            for column in range(1, 13):
                print(
                    str(sheet.cell(1, column).value)
                    + ": "
                    + str(sheet.cell(row, column).value)
                )

            found = True
            break

    if not found:
        print("Employee not found.")


def process_employees(sheet):
    """Process all employee records and return results and report lines."""
    results = []
    report = []

    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    report.append("BEST BUY MEDICAL SUPPLIES INC.")
    report.append("EMPLOYEE PROMOTION REPORT")
    report.append("Generated: " + generated_at)
    report.append("=" * 70)

    for row in range(2, sheet.max_row + 1):
        try:
            employee_id = str(sheet.cell(row, 1).value).strip()
            first_name = str(sheet.cell(row, 2).value).strip()
            last_name = str(sheet.cell(row, 3).value).strip()
            email = str(sheet.cell(row, 4).value).strip()
            department = str(sheet.cell(row, 5).value).strip()
            position = str(sheet.cell(row, 6).value).strip()

            years_text = str(sheet.cell(row, 7).value).lower()
            years = int(years_text.replace("years", "").strip())

            score = int(sheet.cell(row, 8).value)
            attendance = int(sheet.cell(row, 9).value)
            warnings = int(sheet.cell(row, 10).value)
            training = str(sheet.cell(row, 11).value).strip()
            current_salary = float(sheet.cell(row, 12).value)

            full_name = first_name + " " + last_name

            if not validate_employee_id(employee_id):
                report.append(employee_id + " - Invalid Employee ID")
                continue

            if not validate_email(email):
                report.append(employee_id + " - Invalid Email")
                continue

            promotion_status = check_promotion(
                years,
                score,
                attendance,
                warnings,
                training,
            )

            id_number = extract_id_number(employee_id)
            market_salary = get_market_salary(position)

            if current_salary < market_salary:
                salary_status = "Below Market"
            elif current_salary > market_salary:
                salary_status = "Above Market"
            else:
                salary_status = "At Market"

            results.append(
                [
                    employee_id,
                    full_name,
                    department,
                    position,
                    years,
                    score,
                    attendance,
                    warnings,
                    training,
                    current_salary,
                    market_salary,
                    salary_status,
                    promotion_status,
                    id_number,
                ]
            )

            report.append(
                employee_id
                + " | "
                + full_name
                + " | "
                + promotion_status
                + " | "
                + salary_status
            )

        except (TypeError, ValueError) as error:
            report.append(
                "Error in row "
                + str(row)
                + ": "
                + str(error)
            )

    return results, report


def save_excel_report(results):
    """Save processed employee results to an Excel workbook."""
    REPORTS_DIR.mkdir(exist_ok=True)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Promotion Results"

    headers = [
        "Employee ID",
        "Full Name",
        "Department",
        "Position",
        "Years Service",
        "Performance Score",
        "Attendance",
        "Warnings",
        "Training Status",
        "Current Salary",
        "Market Salary",
        "Salary Status",
        "Promotion Status",
        "ID Number",
    ]

    sheet.append(headers)

    for row in results:
        sheet.append(row)

    for column in sheet.columns:
        max_length = 0

        for cell in column:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        sheet.column_dimensions[column[0].column_letter].width = (
            max_length + 3
        )

    output_file = REPORTS_DIR / "PromotionResults.xlsx"
    workbook.save(output_file)

    return output_file


def save_text_report(report, results):
    """Save a text summary report."""
    REPORTS_DIR.mkdir(exist_ok=True)

    eligible_count = 0
    not_eligible_count = 0

    for employee in results:
        if employee[12] == "Eligible":
            eligible_count += 1
        else:
            not_eligible_count += 1

    output_file = REPORTS_DIR / "PromotionReport.txt"

    with open(output_file, "w", encoding="utf-8") as file:
        for line in report:
            file.write(line + "\n")

        file.write("\nPROMOTION SUMMARY\n")
        file.write("Total Employees Processed: " + str(len(results)) + "\n")
        file.write(
            "Employees Eligible for Promotion: "
            + str(eligible_count)
            + "\n"
        )
        file.write(
            "Employees Not Eligible for Promotion: "
            + str(not_eligible_count)
            + "\n"
        )

    return output_file


def save_word_report(report, results):
    """Save a formatted Word report."""
    REPORTS_DIR.mkdir(exist_ok=True)

    eligible_count = 0
    not_eligible_count = 0

    for employee in results:
        if employee[12] == "Eligible":
            eligible_count += 1
        else:
            not_eligible_count += 1

    document = Document()
    document.add_heading("BEST BUY MEDICAL SUPPLIES INC.", level=1)
    document.add_heading(
        "Employee Promotion Eligibility Report",
        level=2,
    )

    document.add_paragraph(
        "Generated on: "
        + datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    )

    document.add_heading("Employee Results", level=2)

    for line in report:
        document.add_paragraph(line)

    document.add_heading("Promotion Summary", level=2)
    document.add_paragraph(
        "Total Employees Processed: " + str(len(results))
    )
    document.add_paragraph(
        "Employees Eligible for Promotion: " + str(eligible_count)
    )
    document.add_paragraph(
        "Employees Not Eligible for Promotion: "
        + str(not_eligible_count)
    )

    output_file = REPORTS_DIR / "PromotionReport.docx"
    document.save(output_file)

    return output_file


def process_promotion_report():
    """Run the complete reporting workflow."""
    sheet = load_employee_sheet()
    results, report = process_employees(sheet)

    excel_file = save_excel_report(results)
    text_file = save_text_report(report, results)
    word_file = save_word_report(report, results)

    print("\nReports generated successfully!")
    print("Excel: " + str(excel_file))
    print("Text:  " + str(text_file))
    print("Word:  " + str(word_file))


def menu():
    """Display the application menu."""
    while True:
        print("\n" + "=" * 45)
        print("EMPLOYEE PROMOTION & SALARY ANALYSIS")
        print("=" * 45)
        print("1. Display Employees")
        print("2. Search Employee")
        print("3. Process Promotion Report")
        print("4. Exit")

        choice = input("Select option (1-4): ").strip()

        if choice == "1":
            sheet = load_employee_sheet()
            display_employees(sheet)
        elif choice == "2":
            sheet = load_employee_sheet()
            search_employee(sheet)
        elif choice == "3":
            process_promotion_report()
        elif choice == "4":
            print("Program closed.")
            break
        else:
            print("Please select a number from 1 to 4.")


if __name__ == "__main__":
    menu()
